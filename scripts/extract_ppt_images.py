import os
import pptx
from pptx import Presentation
from PIL import Image

folder = r"C:\Users\Albert\Desktop\Antigravity\esmotech\PPT Esmotech\PPT BO"
ppt_file = [f for f in os.listdir(folder) if f.endswith(".pptx")][0]
prs = Presentation(os.path.join(folder, ppt_file))

out_dir = r"C:\Users\Albert\Desktop\Antigravity\esmotech\scratch\extracted_ppt_images"
os.makedirs(out_dir, exist_ok=True)

for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            ext = img.ext
            filename = f"slide_{i+1}_{shape.name}.{ext}".replace(";", "_")
            filepath = os.path.join(out_dir, filename)
            with open(filepath, "wb") as f:
                f.write(img.blob)
            print(f"Slide {i+1}: Saved {filename}, size in PPT: {shape.width.inches:.2f}x{shape.height.inches:.2f} in, dimensions: {img.size}")

# Also inspect the target logo
target_logo = r"C:\Users\Albert\Desktop\Antigravity\esmotech\LOGOS\ESMOTECH_NEGRE_LOGO-upscaled.png"
with Image.open(target_logo) as im:
    print(f"\nTarget logo: {target_logo}")
    print(f"Format: {im.format}, Size: {im.size}, Mode: {im.mode}")
