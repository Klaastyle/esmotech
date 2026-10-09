from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Inicialitzar presentació (Format panoràmic 16:9)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors oficials Esmotech
BG_COLOR = RGBColor(248, 249, 250)      # Fons clar #F8F9FA
CARD_BG = RGBColor(255, 255, 255)       # Targeta blanca #FFFFFF
TEXT_DARK = RGBColor(15, 17, 21)        # Negre fosc #0F1115
BLUE_ACCENT = RGBColor(0, 136, 255)     # Blau elèctric #0088FF
GRAY_TEXT = RGBColor(100, 116, 139)     # Gris acer #64748B

def aplicar_fons_i_peu(slide, numero_diapo):
    # Fons general
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    # Peu de pàgina estandard
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(10), Inches(0.3))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "DOSSIER CORPORATIU I COMERCIAL 2026  |  ESMOTECH CUTTING SOLUTIONS"
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY_TEXT
    
    # Número de diapo a la dreta
    txBox_num = slide.shapes.add_textbox(Inches(11.5), Inches(7.0), Inches(1.0), Inches(0.3))
    tf_num = txBox_num.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = f"{numero_diapo:02d}"
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = BLUE_ACCENT
    p_num.alignment = PP_ALIGN.RIGHT

# ----------------------------------------------------
# DIAPOSITIVA 1: PORTADA
# ----------------------------------------------------
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
aplicar_fons_i_peu(slide1, 1)

# Píndola superior
shape1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.6), Inches(1.0), Inches(6.1), Inches(0.6))
shape1.fill.solid()
shape1.fill.fore_color.rgb = CARD_BG
shape1.line.color.rgb = BLUE_ACCENT
tf1 = shape1.text_frame
p1 = tf1.paragraphs[0]
p1.text = "ESPECIALISTES EN EINES DE TALL I ESMOLAT INDUSTRIAL"
p1.font.size = Pt(13)
p1.font.bold = True
p1.font.color.rgb = BLUE_ACCENT
p1.alignment = PP_ALIGN.CENTER

# Títol Principal
tx_titol = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.2))
tf_t = tx_titol.text_frame
tf_t.word_wrap = True
pt = tf_t.paragraphs[0]
pt.text = "TECNOLOGIA AVANÇADA, MÀXIMA PRECISIÓ I SERVEI TÈCNIC INTEGRAL"
pt.font.size = Pt(26)
pt.font.bold = True
pt.font.color.rgb = TEXT_DARK
pt.alignment = PP_ALIGN.CENTER

# Subtítol / Descripció
tx_sub = slide1.shapes.add_textbox(Inches(2.0), Inches(3.6), Inches(9.3), Inches(0.8))
tf_s = tx_sub.text_frame
tf_s.word_wrap = True
ps = tf_s.paragraphs[0]
ps.text = "Optimitzem el rendiment, reduim mermes i garantim la durabilitat en la indústria alimentària."
ps.font.size = Pt(16)
ps.font.color.rgb = GRAY_TEXT
ps.alignment = PP_ALIGN.CENTER

# 4 Botons / Píndoles inferiors
valors = ["TECNOLOGIA", "PRECISIÓ", "QUALITAT", "SERVEI"]
for i, val in enumerate(valors):
    left = Inches(1.8 + (i * 2.5))
    btn = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(5.0), Inches(2.2), Inches(0.9))
    btn.fill.solid()
    btn.fill.fore_color.rgb = CARD_BG
    btn.line.color.rgb = GRAY_TEXT
    tf_b = btn.text_frame
    pb = tf_b.paragraphs[0]
    pb.text = val
    pb.font.size = Pt(14)
    pb.font.bold = True
    pb.font.color.rgb = TEXT_DARK
    pb.alignment = PP_ALIGN.CENTER


# ----------------------------------------------------
# DIAPOSITIVA 2: PRESENTACIÓ DE L'EMPRESA (3 columnes)
# ----------------------------------------------------
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
aplicar_fons_i_peu(slide2, 2)

# Capçalera
tx_h2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
tf_h2 = tx_h2.text_frame
p_cat = tf_h2.paragraphs[0]
p_cat.text = "PRESENTACIÓ DE L'EMPRESA"
p_cat.font.size = Pt(12)
p_cat.font.bold = True
p_cat.font.color.rgb = BLUE_ACCENT
p_tit = tf_h2.add_paragraph()
p_tit.text = "QUI SOM, QUÈ FEM I COM TREBALLEM"
p_tit.font.size = Pt(22)
p_tit.font.bold = True
p_tit.font.color.rgb = TEXT_DARK

# 3 Columnes (Targetes)
columnes = [
    ("01 / COMPROMÍS — QUI SOM", "Empresa de nova creació recolzada per una dilatada experiència professional en el sector alimentari. Entenem les particularitats i necessitats d'un sector essencial i estratègic, convertint-nos en una garantia de confiança sòlida per a tots els nostres clients.", "[Placeholder: imatge diapo 2 - 1]"),
    ("02 / CAPACITAT — QUÈ FEM", "Oferim servei tècnic especialitzat, subministrament d'eines de tall d'alt rendiment i esmolat professional de tot tipus de ganivetes industrials. Disposem d'un centre d'esmolats propi amb maquinària d'última generació.", "[Placeholder: imatge diapo 2 - 2]"),
    ("03 / GARANTIA — COM HO FEM", "Garantim la màxima qualitat de tall per reduir mermes de producte i prolongar la vida útil de cada eina. Oferim assessorament tècnic continu i mantenim estocs pactats per assegurar un servei àgil.", "[Placeholder: imatge diapo 2 - 3]")
]

for i, (tit, desc, img_ph) in enumerate(columnes):
    left = Inches(0.8 + (i * 4.0))
    # Targeta blanca
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = GRAY_TEXT
    
    tf_c = card.text_frame
    tf_c.word_wrap = True
    p_ct = tf_c.paragraphs[0]
    p_ct.text = tit
    p_ct.font.size = Pt(13)
    p_ct.font.bold = True
    p_ct.font.color.rgb = BLUE_ACCENT
    
    p_cd = tf_c.add_paragraph()
    p_cd.text = "\n" + desc
    p_cd.font.size = Pt(12)
    p_cd.font.color.rgb = TEXT_DARK
    
    p_img = tf_c.add_paragraph()
    p_img.text = f"\n{img_ph}"
    p_img.font.size = Pt(11)
    p_img.font.italic = True
    p_img.font.color.rgb = GRAY_TEXT
    p_img.alignment = PP_ALIGN.CENTER


# ----------------------------------------------------
# LLISTAT DE RESTA DE DIAPOSITIVES ESTRUCTURADES (3 A 12)
# ----------------------------------------------------
restants = [
    (3, "SERVEIS D'ESMOLATS I SOLUCIONS TÈCNIQUES", "Especialitzats per a la indústria alimentària: màxima precisió, eficiència i rendiment constant.\n• Serveis d'esmolats especialitzats\n• Importància de la precisió en el tall\n• Personalització adaptada a cada client\n[Placeholder: imatge diapo 3 - 1]"),
    (4, "TIPUS DE GANIVETES I CENTRES DE TALL", "Gamma completa per a tota la línia de producció:\n• Ganivets / Tisores | Talla-fiambres | Plaques / Ganivetes\n• Culanes | Cutters | Serres circulars | Txuleteres | Llescadores\n[Placeholder: imatge diapo 4 - 1: Placa de picadora]"),
    (5, "SERRES CINTA", "Serres cinta industrials d'alta precisió:\n• Podem fabricar segons mida\n• Stock permanent de referències habituals\n• Materials d'alta qualitat i resistència\n[Placeholder: imatge diapo 5 - 1]"),
    (6, "PRODUCCIÓ I SUBMINISTRAMENT DE GANIVETES", "Solucions estàndard i a mida:\n• Productes Estàndard (Estoc Permanent): Serra cinta, discos, rectes.\n• Productes Especials (Solucions a Mida): Skinners, microdentades, Renxingblade (https://www.renxingblade.com/).\n[Placeholder: imatge diapo 6 - 1]"),
    (7, "CUTTERS I SKINNERS", "Eines dissenyades per a processos d'alta exigència mecànica i sanitària.\n[Placeholder: imatge diapo 7 - 1]"),
    (8, "CAPÇALS TÈRMICS I GANIVETES D'ENVASAT", "Solucions de tall per a línies de termoformat i envasadores d'alta cadència.\n[Placeholder: imatge diapo 8 - 1]"),
    (9, "ALIANÇA HAGEDORN: GANIVETES EVOLUTIVES I CIRCULARS", "Distribuïdor oficial Hagedorn Foodtech (Made in Germany):\n• Compatibilitat universal amb totes les marques del mercat.\n• Rendiment extraordinari i vida útil perllongada.\n[Placeholder: imatge diapo 9 - 1]"),
    (10, "GANIVETES EVOLUTIVES I SERVEI DE REGENERACIÓ", "Tecnologia de materials i regeneració:\n• TRITECH (Alta higiene) | LONGLIFE (Càrregues severes) | GENESIS (Màxima vida)\n• Procés exclusiu de regeneració i trempat al 100% del cos.\n[Placeholder: imatge diapo 10 - 1, 2, 3, 4]"),
    (11, "CATÀLEG AUXILIAR I SERVEIS TÈCNICS", "Subministrament i assistència integral:\n• Capçals tèrmics (Bizerba, Zebra, etc.) | Serres cinta | Reparacions motors Bizerba\n• Recanvis especials (Wizzard, Ibex) | Rediamantat de moles | Abraisus\n[Placeholder: imatge diapo 11 - 1]"),
    (12, "EL VOSTRE PARTNER TÈCNIC DE CONFIANÇA", "A ESMOTECH combinem l'experiència tècnica, la tecnologia de tall d'última generació i la màxima proximitat per garantir la tranquil·litat de la vostra planta de producció.")
]

for num, titol, descripcio in restants:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    aplicar_fons_i_peu(slide, num)
    
    # Capçalera
    tx_h = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
    tf_h = tx_h.text_frame
    p_c = tf_h.paragraphs[0]
    p_c.text = "SOLUCIONS I SERVEIS ESMOTECH"
    p_c.font.size = Pt(12)
    p_c.font.bold = True
    p_c.font.color.rgb = BLUE_ACCENT
    p_t = tf_h.add_paragraph()
    p_t.text = titol
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_DARK
    
    # Targeta de contingut central
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = GRAY_TEXT
    
    tf_body = card.text_frame
    tf_body.word_wrap = True
    p_b = tf_body.paragraphs[0]
    p_b.text = descripcio
    p_b.font.size = Pt(14)
    p_b.font.color.rgb = TEXT_DARK

# Guardar fitxer final
prs.save("Esmotech_Presentacio_2026.pptx")
print("S'ha generat correctament el fitxer: Esmotech_Presentacio_2026.pptx")
