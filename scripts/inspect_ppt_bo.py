import os
import pptx
from pptx import Presentation

folder = r"C:\Users\Albert\Desktop\Antigravity\esmotech\PPT Esmotech\PPT BO"
files = os.listdir(folder)
print("Files in folder:", files)

ppt_path = os.path.join(folder, files[0])
print("Opening:", ppt_path)
prs = Presentation(ppt_path)
print(f"Slide width: {prs.slide_width.inches}, height: {prs.slide_height.inches}")
print(f"Number of slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    texts = []
    images = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    texts.append(txt)
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            images.append(f"{s.name} ({s.width.inches:.1f}x{s.height.inches:.1f} in)")
    print(f"\n--- Slide {i+1} ---")
    print(f"Images count: {len(images)}", images)
    print("Texts snippet:")
    for t in texts[:6]:
        print("  -", t)
