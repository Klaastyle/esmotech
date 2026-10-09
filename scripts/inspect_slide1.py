import os
import pptx
from pptx import Presentation

folder = r"C:\Users\Albert\Desktop\Antigravity\esmotech\PPT Esmotech\PPT BO"
ppt_file = [f for f in os.listdir(folder) if f.endswith(".pptx")][0]
prs = Presentation(os.path.join(folder, ppt_file))

slide1 = prs.slides[0]
print("=== Slide 1 Shapes ===")
for s in slide1.shapes:
    print(f"Shape: {s.name}, Type: {s.shape_type}, Left: {s.left.inches:.2f}, Top: {s.top.inches:.2f}, Width: {s.width.inches:.2f}, Height: {s.height.inches:.2f}")
    if s.has_text_frame:
        print("  Text:", " ".join(p.text for p in s.text_frame.paragraphs if p.text))
