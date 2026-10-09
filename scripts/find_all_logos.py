import os
import pptx
from pptx import Presentation

folder = r"C:\Users\Albert\Desktop\Antigravity\esmotech\PPT Esmotech\PPT BO"
ppt_file = [f for f in os.listdir(folder) if f.endswith(".pptx")][0]
prs = Presentation(os.path.join(folder, ppt_file))

for i, slide in enumerate(prs.slides):
    logos = []
    for s in slide.shapes:
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            # Check if this shape is one of the logo images
            if s.image.size == (1024, 341):
                logos.append(f"Logo shape: {s.name} at ({s.left.inches:.2f}, {s.top.inches:.2f}), size ({s.width.inches:.2f}x{s.height.inches:.2f})")
    print(f"Slide {i+1}: {len(logos)} logos found -> {logos}")
